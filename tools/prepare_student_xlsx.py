from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\arhre\Downloads\student-list-2026.xlsx")
OUTPUT = Path(r"D:\ims\import_data\student-list-2026.json")

SHEET_MAP = {
    "class6-2026": {"prefix": "2026-C6A", "class_id": 2, "section_id": 3, "label": "Class 6 / A"},
    "class7-2026": {"prefix": "2026-C7S1", "class_id": 3, "section_id": 2, "label": "Class 7 / 1"},
    "class8-2026": {"prefix": "2026-C8S1", "class_id": 4, "section_id": 2, "label": "Class 8 / 1"},
    "class9-2026-science": {"prefix": "2026-C9SCI", "class_id": 5, "section_id": 5, "label": "Class 9 / Science"},
    "class9-2026-commerce": {"prefix": "2026-C9COM", "class_id": 5, "section_id": 6, "label": "Class 9 / Commerce"},
    "class9-2026-arts": {"prefix": "2026-C9HUM", "class_id": 5, "section_id": 7, "label": "Class 9 / Humanities"},
    "class10-2026-science": {"prefix": "2026-C10SCI", "class_id": 6, "section_id": 5, "label": "Class 10 / Science"},
    "class10-2026-commerce": {"prefix": "2026-C10COM", "class_id": 6, "section_id": 6, "label": "Class 10 / Commerce"},
    "class10-2026-arts": {"prefix": "2026-C10HUM", "class_id": 6, "section_id": 7, "label": "Class 10 / Humanities"},
}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def normalize_gender(value: str) -> str:
    value = value.strip().lower()
    if value in {"নারী", "female", "f"}:
        return "Female"
    if value in {"পুরুষ", "male", "m"}:
        return "Male"
    return value


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    output = []
    summary = []

    for sheet_name, mapping in SHEET_MAP.items():
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(cell) for cell in next(rows)]
        count = 0

        for raw_row in rows:
            row = dict(zip(headers, raw_row))
            if not any(clean(value) for value in row.values()):
                continue

            count += 1
            record = {key: clean(row.get(key)) for key in headers}
            if not record.get("first_name") and record.get("middlename"):
                record["first_name"] = record["middlename"]
                record["middlename"] = ""
            if not (record.get("first_name") or record.get("middlename") or record.get("last_name")):
                count -= 1
                continue
            record["gender"] = normalize_gender(record.get("gender", ""))
            record["admission_no"] = f"{mapping['prefix']}-{count:03d}"
            record["original_admission_no"] = clean(row.get("admission_no"))
            record["class_id"] = mapping["class_id"]
            record["section_id"] = mapping["section_id"]
            record["sheet"] = sheet_name
            record["label"] = mapping["label"]
            output.append(record)

        summary.append({"sheet": sheet_name, "label": mapping["label"], "rows": count})

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({"summary": summary, "students": output}, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Wrote {len(output)} students to {OUTPUT}")
    for item in summary:
        print(f"{item['label']} ({item['sheet']}): {item['rows']}")


if __name__ == "__main__":
    main()
